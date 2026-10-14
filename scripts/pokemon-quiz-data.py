"""Build sites/pokemon-quiz/quiz-data.js from the quiz spreadsheet, and
download any sprites it needs that aren't in sites/pokemon-quiz/sprites/ yet.

Usage (needs `pip install openpyxl pillow`):
  python3 scripts/pokemon-quiz-data.py Pokemon_Personality_Quiz.xlsx sites/pokemon-quiz

Sprites are Pokémon Showdown's static gen 5 set
(https://play.pokemonshowdown.com/sprites/gen5/), except Megas, which use
Serebii's official artwork (https://www.serebii.net/pokemon/megaevolution.shtml),
shrunk to WebP in sprites/mega/. Prints the results that have no exact sprite
(they fall back to the regular form). See sites/pokemon-quiz/README.md.
"""
import json, os, re, subprocess, sys, unicodedata
import openpyxl
from PIL import Image

SPRITE_URL = 'https://play.pokemonshowdown.com/sprites/gen5/'
SEREBII = 'https://www.serebii.net'
MEGA_ART_PX = 384  # shown at 192px, so 2x for retina screens
xlsx, out_dir = sys.argv[1:3]

def fetch(url, out=None):
    # curl rather than urllib: some Python installs ship without CA certificates
    args = ['curl', '-sfL', '-A', 'Mozilla/5.0', url] + (['-o', out] if out else [])
    res = subprocess.run(args, check=True, capture_output=True)
    return None if out else res.stdout.decode('latin-1')

listing = fetch(SPRITE_URL + '?view=dir')
SPRITES = set(re.findall(r'href="(?:[^"]*/)?([a-z0-9-]+)\.png"', listing))

def sid(s):
    """Showdown's id: lowercase letters and digits only (Mr. Mime -> mrmime)."""
    s = s.replace('♀', 'f').replace('♂', 'm')
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]', '', s.lower())

REGION = {'Alolan': 'alola', 'Galarian': 'galar', 'Hisuian': 'hisui', 'Paldean': 'paldea'}
SPECIAL = {
    'Paldean Tauros (Combat Breed)': 'tauros-paldeacombat',
    'Paldean Tauros (Blaze Breed)': 'tauros-paldeablaze',
    'Paldean Tauros (Aqua Breed)': 'tauros-paldeaaqua',
    'Hisuian Basculin (White-Striped)': 'basculin-whitestriped',
    'Galarian Darmanitan (Zen Mode)': 'darmanitan-galarzen',
    'Ultra Necrozma': 'necrozma-ultra',
    'Gigantamax Urshifu (Rapid Strike)': 'urshifu-rapidstrikegmax',
    'Gigantamax Urshifu (Single Strike)': 'urshifu-gmax',
    'Mega Meowstic': 'meowstic-mmega',
}

def species_of(name):
    n = re.sub(r'\s*\(.*\)$', '', name)
    n = re.sub(r'^(Mega|Gigantamax|Primal|Ultra|Alolan|Galarian|Hisuian|Paldean|Fan|Frost|Heat|Mow|Wash)\s+', '', n)
    return re.sub(r'\s+[XYZ]$', '', n)

def candidate(name, form):
    if name in SPECIAL: return SPECIAL[name]
    m = re.match(r'^(.*) \((\w+) Forme\)$', name)
    if m: return sid(m.group(1)) + '-' + sid(m.group(2))
    first, rest = name.split(' ', 1) if ' ' in name else (name, '')
    if form == 'Base': return sid(name)
    if form == 'Regional': return sid(rest) + '-' + REGION[first]
    if form == 'Alternate' and first in ('Fan', 'Frost', 'Heat', 'Mow', 'Wash'): return 'rotom-' + first.lower()
    if form == 'Gigantamax': return sid(rest) + '-gmax'
    if form == 'Primal': return sid(rest) + '-primal'
    if form == 'Mega':
        m = re.match(r'^Mega (.*?)(?: ([XYZ]))?$', name)
        return sid(m.group(1)) + '-mega' + (m.group(2) or '').lower() if m else None
    return None

def sprite_for(name, form):
    """Returns (file, exact)."""
    c = candidate(name, form)
    if c in SPRITES: return c, True
    m = re.match(r'^Mega (.*) \((\w+)\)$', name)  # Mega Magearna (Original) -> magearna-original
    if m and sid(m.group(1)) + '-' + sid(m.group(2)) in SPRITES:
        return sid(m.group(1)) + '-' + sid(m.group(2)), False
    base = sid(species_of(name))
    return (base if base in SPRITES else None), False

def slug(s):
    """File-name slug, matching the share-link ids in index.html."""
    s = s.replace('♀', '-f').replace('♂', '-m')
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '-', re.sub(r"[’'.:()]", '', s)).strip('-')

def mega_art_urls():
    """{quiz result name: full-size artwork URL} from Serebii's Mega Evolution list."""
    page = fetch(SEREBII + '/pokemon/megaevolution.shtml')
    art = {}
    for attrs, label in re.findall(r'<img ([^>]*class="pkmn"[^>]*)>\s*</a></td></tr>\s*'
                                   r'<tr><td align="center"><a [^>]*>(.*?)</a>', page, re.S):
        src = re.search(r'src="([^"]+)"', attrs).group(1).replace('/art/small/', '/art/')
        name, _, variant = re.sub(r'<br\s*/?>', '|', label).partition('|')
        name = name.strip()
        if variant.startswith('Original'): name += ' (Original)'
        # Tatsugiri: use the Stretchy Form art for all three Mega Tatsugiri results
        if name == 'Mega Tatsugiri' and 'Stretchy' not in variant: continue
        art[name] = SEREBII + src
    return art

def mega_key(name):
    return name if name.endswith('(Original)') else re.sub(r'\s*\(.*\)$', '', name)

MEGA_ART = mega_art_urls()

wb = openpyxl.load_workbook(xlsx)
qrows = list(wb['Question tree'].iter_rows(min_row=2, values_only=True))
rrows = list(wb['Results'].iter_rows(min_row=2, values_only=True))

results, rindex = [], {}
mega_files = {}  # sprite path -> artwork URL
missing = []
for name, form, dex, typ, nq, path, vibe in rrows:
    if not name: continue
    if form == 'Mega' and mega_key(name) in MEGA_ART:
        f, exact = f'mega/{slug(name)}.webp', True
        mega_files[f] = MEGA_ART[mega_key(name)]
    else:
        f, exact = sprite_for(name, form)
        if f: f += '.png'
    if not exact: missing.append((name, form, f))
    rindex[name] = len(results)
    results.append({'n': name, 'f': form, 'd': dex, 't': typ, 'q': nq, 'v': vibe, 's': f, 'x': exact})

def target(t):
    t = str(t).strip()
    if re.fullmatch(r'Q\d+', t): return int(t[1:])
    m = re.match(r'^Result:\s*(.*)$', t)
    name = m.group(1).strip() if m else t
    if name not in rindex: raise SystemExit(f'unknown target {t!r}')
    return -(rindex[name] + 1)  # negative = result index

questions = {}
for qn, lvl, text, a, ga, b, gb, trait in qrows:
    if qn is None: continue
    questions[int(qn)] = [text, a, target(ga), b, target(gb)]

# sanity: every result reachable exactly once, every question reachable
reached_r, reached_q, stack = [], set(), [1]
while stack:
    q = stack.pop(); reached_q.add(q)
    for t in (questions[q][2], questions[q][4]):
        (stack.append(t) if t > 0 else reached_r.append(-t - 1))
assert len(reached_q) == len(questions), (len(reached_q), len(questions))
assert sorted(reached_r) == list(range(len(results))), 'results not reached exactly once'

qlist = [questions[i] for i in range(1, len(questions) + 1)]
needed = sorted({r['s'] for r in results if r['s']})
with open(f'{out_dir}/quiz-data.js', 'w') as fh:
    fh.write('// Generated from Pokemon_Personality_Quiz.xlsx — do not edit by hand.\n')
    fh.write('// Q[i] = question i+1: [text, optionA, nextA, optionB, nextB];\n')
    fh.write('// next > 0 is a question number, next < 0 is result index (-next - 1).\n')
    fh.write('// R[i]: n name, f form, d dex no., t type, q questions, v vibe, s sprite path, x sprite is exact.\n')
    fh.write('const Q = ' + json.dumps(qlist, ensure_ascii=False, separators=(',', ':')) + ';\n')
    fh.write('const R = ' + json.dumps(results, ensure_ascii=False, separators=(',', ':')) + ';\n')
os.makedirs(f'{out_dir}/sprites/mega', exist_ok=True)
to_fetch = [f for f in needed if not os.path.exists(f'{out_dir}/sprites/{f}')]
for f in to_fetch:
    dest = f'{out_dir}/sprites/{f}'
    if f in mega_files:
        fetch(mega_files[f], dest + '.tmp')
        im = Image.open(dest + '.tmp')
        im.thumbnail((MEGA_ART_PX, MEGA_ART_PX), Image.LANCZOS)
        im.save(dest, 'WEBP', quality=88, method=6)
        os.remove(dest + '.tmp')
    else:
        fetch(SPRITE_URL + f, dest)
print(len(qlist), 'questions,', len(results), 'results,', len(needed), 'sprites,', len(missing), 'without exact sprite')
for name, form, fallback in missing:
    print(f'  no exact sprite: {name} ({form}) -> {fallback or "placeholder"}')
print(f'downloaded {len(to_fetch)} new sprite(s)')
on_disk = {os.path.relpath(os.path.join(d, f), f'{out_dir}/sprites')
           for d, _, fs in os.walk(f'{out_dir}/sprites') for f in fs}
unused = sorted(on_disk - set(needed))
if unused: print('unused sprite files (safe to delete):', ' '.join(unused))
